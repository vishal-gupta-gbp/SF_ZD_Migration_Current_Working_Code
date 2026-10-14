"""Seeds field and value mapping tables from the mapping workbook into the database."""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from sqlalchemy.orm import Session

from migration.db.models import MappingField, MappingValue


def s(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip()


def normalize_label(val: Any) -> str:
    return s(val).strip().lower()


def seed_mappings_if_needed(session: Session, mapping_path: Path, force: bool = False) -> Dict[str, int]:
    """Seed mapping_fields and mapping_values from XLSX into DB if empty or if force=True."""
    field_count = session.query(MappingField).count()
    value_count = session.query(MappingValue).count()

    if not force and field_count > 0 and value_count > 0:
        return {"field_mappings": field_count, "value_mappings": value_count, "seeded": False}

    if not mapping_path.exists():
        print(f"Warning: Mapping file {mapping_path} not found. Skipping mapping seed.")
        return {"field_mappings": field_count, "value_mappings": value_count, "seeded": False}

    print(f"Loading mapping definitions from {mapping_path.name} into database...", flush=True)

    if force:
        session.query(MappingField).delete()
        session.query(MappingValue).delete()
        session.flush()

    wb = openpyxl.load_workbook(mapping_path, read_only=True, data_only=True)
    fields_to_add: List[MappingField] = []
    values_to_add: List[MappingValue] = []

    try:
        if "SF to ZD Field Map" in wb.sheetnames:
            ws = wb["SF to ZD Field Map"]
            headers = [s(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
            for row in ws.iter_rows(min_row=2, values_only=True):
                item = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
                sf_api = s(item.get("Salesforce API Name"))
                if not sf_api:
                    continue

                raw_fid = s(item.get("Zendesk Target Field ID"))
                target_fid = None
                if raw_fid:
                    try:
                        target_fid = int(float(raw_fid))
                    except (ValueError, TypeError):
                        pass

                fields_to_add.append(
                    MappingField(
                        salesforce_api_name=sf_api,
                        salesforce_label=s(item.get("Salesforce Field Label")),
                        salesforce_type=s(item.get("Salesforce Type")),
                        zendesk_target_field=s(item.get("Zendesk Target Field")),
                        zendesk_target_field_id=target_fid,
                        zendesk_target_type=s(item.get("Zendesk Target Type")),
                        mapping_status=s(item.get("Mapping Status")),
                        transformation_notes=s(item.get("Transformation / Logic")),
                    )
                )

        if "Value Mapping" in wb.sheetnames:
            ws = wb["Value Mapping"]
            headers = [s(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
            for row in ws.iter_rows(min_row=2, values_only=True):
                item = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
                sf_api = s(item.get("Salesforce API Name"))
                sf_val = s(item.get("Salesforce Value"))
                if not sf_api or not sf_val:
                    continue

                values_to_add.append(
                    MappingValue(
                        salesforce_api_name=sf_api,
                        salesforce_value=sf_val,
                        zendesk_value=s(item.get("Zendesk Value")),
                        zendesk_tag=s(item.get("Zendesk Tag")),
                        match_status=s(item.get("Match Status")),
                    )
                )
    finally:
        wb.close()

    if fields_to_add:
        session.bulk_save_objects(fields_to_add)
    if values_to_add:
        session.bulk_save_objects(values_to_add)
    session.commit()

    print(
        f"Database mapping seeded successfully: {len(fields_to_add):,} fields, "
        f"{len(values_to_add):,} values.",
        flush=True,
    )
    return {
        "field_mappings": len(fields_to_add),
        "value_mappings": len(values_to_add),
        "seeded": True,
    }
