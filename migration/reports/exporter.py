"""Export comprehensive reconciliation reports directly from database to Excel/CSV."""
from __future__ import annotations

import csv
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from migration.db.models import (
    MigrationAttachment,
    MigrationTicket,
    ZendeskOrganization,
    ZendeskUser,
)


def export_migration_reports(session: Session, output_dir: Path) -> Dict[str, str]:
    """Export reconciliation Excel files and CSVs directly from database tables."""
    output_dir.mkdir(parents=True, exist_ok=True)
    summary_xlsx = output_dir / "Migration_Summary.xlsx"

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)

    # -------------------------------------------------------------
    # 1. Sheet: Executive Summary
    # -------------------------------------------------------------
    ws_sum = wb.create_sheet(title="Executive Summary")
    ws_sum.append(["Salesforce to Zendesk Migration - Reconciliation Summary"])
    ws_sum.append(["Generated at (UTC):", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")])
    ws_sum.append([])

    tickets = session.query(MigrationTicket).all()
    total_tickets = len(tickets)
    completed = sum(1 for t in tickets if t.status == "completed")
    completed_warn = sum(1 for t in tickets if t.status == "completed_with_warnings")
    existing = sum(1 for t in tickets if t.status == "existing")
    failed = sum(1 for t in tickets if t.status == "failed")
    review = sum(1 for t in tickets if t.status == "manual_review_required")
    standard_count = sum(1 for t in tickets if t.ticket_type == "standard")
    phone_count = sum(1 for t in tickets if t.ticket_type == "phone")

    orgs = session.query(ZendeskOrganization).all()
    users = session.query(ZendeskUser).all()
    attachments = session.query(MigrationAttachment).all()
    att_downloaded = sum(1 for a in attachments if a.download_status == "success")
    att_uploaded = sum(1 for a in attachments if a.upload_status == "success")

    summary_rows = [
        ["Total Salesforce Cases Evaluated", total_tickets],
        ["Standard Cases", standard_count],
        ["Phone Origin Cases", phone_count],
        ["Imported Tickets (Success)", completed],
        ["Imported with Warnings (Attachments)", completed_warn],
        ["Already Existing in Zendesk", existing],
        ["Failed Ticket Imports", failed],
        ["Manual Review Required", review],
        ["Total Zendesk Organizations Synced", len(orgs)],
        ["Total Zendesk End Users Synced", len(users)],
        ["S3 Attachments Downloaded", att_downloaded],
        ["S3 Attachments Uploaded to Zendesk", att_uploaded],
    ]

    ws_sum.append(["Metric", "Count"])
    for cell in ws_sum[4]:
        cell.fill = header_fill
        cell.font = header_font

    for row in summary_rows:
        ws_sum.append(row)

    # -------------------------------------------------------------
    # 2. Sheet: Ticket Details
    # -------------------------------------------------------------
    ws_details = wb.create_sheet(title="Ticket Reconciliation")
    detail_headers = [
        "Salesforce Case ID", "Case Number", "Type", "Status", "Zendesk Ticket ID",
        "Zendesk External ID", "External ID Match", "Attempts", "HTTP Status",
        "Error Type", "Failure Reason", "Started At", "Completed At"
    ]
    ws_details.append(detail_headers)
    for cell in ws_details[1]:
        cell.fill = header_fill
        cell.font = header_font

    for t in tickets:
        id_match = "YES" if t.zendesk_external_id and t.zendesk_external_id == t.case_id else ("NO" if t.zendesk_ticket_id else "")
        ws_details.append([
            t.case_id,
            t.case_number or "",
            t.ticket_type or "",
            t.status or "",
            t.zendesk_ticket_id or "",
            t.zendesk_external_id or "",
            id_match,
            t.attempts or 0,
            t.http_status or "",
            t.error_type or "",
            t.failure_reason or "",
            t.started_at.strftime("%Y-%m-%d %H:%M:%S") if t.started_at else "",
            t.completed_at.strftime("%Y-%m-%d %H:%M:%S") if t.completed_at else "",
        ])

    # -------------------------------------------------------------
    # 3. Sheet: Attachments
    # -------------------------------------------------------------
    ws_att = wb.create_sheet(title="Attachment Reconciliation")
    att_headers = [
        "Case ID", "Comment ID", "File Name", "Download Status", "Upload Status",
        "File Size Bytes", "SHA256", "Local Path", "Upload Token", "Download Error", "Upload Error"
    ]
    ws_att.append(att_headers)
    for cell in ws_att[1]:
        cell.fill = header_fill
        cell.font = header_font

    for a in attachments:
        ws_att.append([
            a.case_id,
            a.comment_id or "",
            a.file_name or "",
            a.download_status or "",
            a.upload_status or "",
            a.file_size_bytes or 0,
            a.sha256 or "",
            a.local_file_path or "",
            a.upload_token or "",
            a.download_error or "",
            a.upload_error or "",
        ])

    # Auto-adjust column widths
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 50)

    wb.save(summary_xlsx)
    print(f"Reconciliation workbook exported from database: {summary_xlsx}", flush=True)

    # Also export CSV for quick CLI consumption
    ticket_csv = output_dir / "Migration_Tickets.csv"
    with open(ticket_csv, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(detail_headers)
        for t in tickets:
            id_match = "YES" if t.zendesk_external_id and t.zendesk_external_id == t.case_id else ""
            writer.writerow([
                t.case_id, t.case_number or "", t.ticket_type or "", t.status or "",
                t.zendesk_ticket_id or "", t.zendesk_external_id or "", id_match,
                t.attempts or 0, t.http_status or "", t.error_type or "", t.failure_reason or "",
                t.started_at or "", t.completed_at or ""
            ])

    return {
        "excel_summary": str(summary_xlsx),
        "tickets_csv": str(ticket_csv),
    }
